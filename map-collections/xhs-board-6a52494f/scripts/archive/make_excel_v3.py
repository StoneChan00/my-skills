# -*- coding: utf-8 -*-
"""
基于官方新模板生成导入文件 amap-import-v2.xlsx。
模板列: 名称 | *经度 | *纬度 | *地址 | 颜色 | 图标(外轮廓) | 图标(填充物) | 描述 | 文件夹
规则: 名称≤40字, 描述≤600字, 颜色1-8, 经纬度优先匹配, sheet 名保持「标记位置」不变。
"""
import json
import shutil
import warnings
from pathlib import Path

from openpyxl import load_workbook

warnings.filterwarnings("ignore")

BASE = Path(__file__).resolve().parent.parent
TPL = BASE / "导入模板.xlsx"
OUT = BASE / "amap-import-v2.xlsx"
GEO = BASE / "osm_geocoded_final.json"

# 颜色映射（对照「颜色说明」: 1蓝 2浅蓝 3青 4绿 5黄 6橙 7鲑红 8红）
COLOR_MAP = {"景点": 1, "美食": 6, "按摩": 8, "住宿": 4, "购物": 5, "交通": 3}


def clip(s, n):
    s = str(s or "")
    return s if len(s) <= n else s[: n - 1] + "…"


def main():
    data = [d for d in json.loads(GEO.read_text(encoding="utf-8")) if d.get("status") == "ok"]

    shutil.copy(TPL, OUT)
    wb = load_workbook(OUT)
    ws = wb["标记位置"]

    # 清掉示例行（r2-r5）
    ws.delete_rows(2, ws.max_row - 1)

    city_order = {"胡志明市": 0, "芽庄": 1, "岘港": 2}
    cat_order = {"景点": 0, "美食": 1, "按摩": 2, "住宿": 3, "购物": 4, "交通": 5}
    data.sort(key=lambda d: (city_order.get(d["city"], 9), cat_order.get(d["category"], 9)))

    n = 0
    for d in data:
        addr_parts = ["越南", d["city"]]
        if d.get("address"):
            addr_parts.append(d["address"])
        addr = " · ".join(addr_parts)

        desc = " | ".join(f"《{x}》" for x in d.get("notes", []))
        if d.get("approx"):
            desc += f"｜⚠位置约为{d.get('osm_name', '')}附近，建议手动微调"
        if d.get("address"):
            desc += f"｜{d['address']}"

        ws.append([
            clip(d["display"], 40),      # 名称
            round(float(d["lon"]), 6),   # *经度
            round(float(d["lat"]), 6),   # *纬度
            clip(addr, 120),             # *地址
            COLOR_MAP.get(d.get("category"), 7),  # 颜色
            None,                        # 图标(外轮廓)
            None,                        # 图标(填充物)
            clip(desc, 600),             # 描述
            d["city"],                   # 文件夹
        ])
        n += 1

    wb.save(OUT)
    print(f"生成 {OUT.name}: {n} 行数据（基于官方模板）")

    # 验证读回
    wb2 = load_workbook(OUT)
    ws2 = wb2["标记位置"]
    rows = list(ws2.iter_rows(min_row=1, max_row=5, values_only=True))
    print("表头:", rows[0])
    for r in rows[1:4]:
        print("样例:", str(r[0])[:26], "|", r[1], ",", r[2], "|", r[4], "|", r[8])
    print(f"总行数: {ws2.max_row} (1 表头 + {ws2.max_row - 1} 数据)")
    print(f"sheet 保留: {wb2.sheetnames}")


if __name__ == "__main__":
    main()
