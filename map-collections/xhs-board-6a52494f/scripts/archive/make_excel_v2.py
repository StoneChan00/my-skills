# -*- coding: utf-8 -*-
"""
生成高德「地图小程序」批量导入 Excel（海外版，越南）。
经纬度为 WGS-84（海外无 GCJ-02 偏移，可直接导入）。
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE = Path(__file__).resolve().parent.parent
IN_FILE = BASE / "osm_geocoded_final.json"
OUT_FILE = BASE / "amap-import.xlsx"

HEADERS = ["名称", "经纬度", "地址", "颜色", "描述", "文件夹"]
COLOR_MAP = {"景点": 1, "美食": 6, "按摩": 8, "住宿": 3, "购物": 5, "交通": 4}
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")


def main():
    data = [d for d in json.loads(IN_FILE.read_text(encoding="utf-8")) if d.get("status") == "ok"]
    print(f"点位: {len(data)}")

    wb = Workbook()
    ws = wb.active
    ws.title = "导入数据"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center")

    city_order = {"胡志明市": 0, "芽庄": 1, "岘港": 2}
    cat_order = {"景点": 0, "美食": 1, "按摩": 2, "住宿": 3, "购物": 4, "交通": 5}
    data.sort(key=lambda d: (city_order.get(d["city"], 9), cat_order.get(d["category"], 9)))

    for d in data:
        addr_parts = ["越南", d["city"]]
        if d.get("address"):
            addr_parts.append(d["address"])
        elif d.get("osm_name") and d.get("precision") not in ("approx",):
            addr_parts.append(str(d["osm_name"]))
        addr = " · ".join(addr_parts)

        desc = " | ".join(f"《{n}》" for n in d.get("notes", []))
        if d.get("approx"):
            hint = d.get("osm_name", "")
            desc += f"｜⚠位置约为{hint}附近，建议在高德中手动微调"
        if d.get("address"):
            desc += f"｜{d['address']}"

        row = [
            d["display"],
            f"{d['lon']:.6f},{d['lat']:.6f}",
            addr,
            COLOR_MAP.get(d.get("category"), 7),
            desc,
            d["city"],
        ]
        ws.append(row)

    for col, w in zip("ABCDEF", (32, 20, 44, 6, 55, 10)):
        ws.column_dimensions[col].width = w
    wb.save(OUT_FILE)

    print(f"生成 {OUT_FILE}")
    stats = {}
    for d in data:
        k = f"{d['city']}/{d['category']}"
        stats[k] = stats.get(k, 0) + 1
    for k in sorted(stats):
        print(f"  {k}: {stats[k]}")
    ap = sum(1 for d in data if d.get("approx"))
    print(f"近似定位: {ap} 个（描述列已标注 ⚠）")


if __name__ == "__main__":
    main()
