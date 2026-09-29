# -*- coding: utf-8 -*-
"""解析高德导入失败回传文件: 找错误标注列/批注/单元格差异。"""
import warnings
from pathlib import Path

from openpyxl import load_workbook

warnings.filterwarnings("ignore")

F = Path(r"D:\2_Projects\github\StoneChan00\my-skills\map-collections\xhs-board-6a52494f\失败文件.xlsx")
print(f"file: {F} ({F.stat().st_size} bytes)")

wb = load_workbook(F)
print(f"sheets: {wb.sheetnames}")

for ws in wb.worksheets:
    print(f"\n===== [{ws.title}] max_row={ws.max_row} max_col={ws.max_column} =====")
    # 批注
    for row in ws.iter_rows():
        for c in row:
            if c.comment:
                print(f"  COMMENT @{c.coordinate}: {c.comment.text[:200]}")
    # 全部行（前10 + 抽样）
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 10)):
        vals = [f"{c.coordinate}:{str(c.value)[:30]}" for c in row if c.value is not None]
        if vals:
            print("  " + " | ".join(vals))
    if ws.max_row > 10:
        for row in ws.iter_rows(min_row=ws.max_row - 2, max_row=ws.max_row):
            vals = [f"{c.coordinate}:{str(c.value)[:30]}" for c in row if c.value is not None]
            if vals:
                print("  " + " | ".join(vals))
        print(f"  ...(共 {ws.max_row} 行)")
