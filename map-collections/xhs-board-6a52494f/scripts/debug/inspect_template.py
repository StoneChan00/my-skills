# -*- coding: utf-8 -*-
"""解剖官方导入模板: sheet 结构、表头、示例行、数据校验、合并单元格。"""
from pathlib import Path

from openpyxl import load_workbook

F = Path(r"D:\2_Projects\github\StoneChan00\my-skills\map-collections\xhs-board-6a52494f\导入模板.xlsx")
print(f"file: {F} ({F.stat().st_size} bytes)")

wb = load_workbook(F)
print(f"sheets: {wb.sheetnames}")

for ws in wb.worksheets:
    print(f"\n===== sheet [{ws.title}] dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column} =====")
    if ws.merged_cells.ranges:
        print("merged:", [str(r) for r in ws.merged_cells.ranges][:6])
    # 数据校验
    for dv in ws.data_validations.dataValidation:
        print(f"  DV: {dv.type} {dv.formula1} @ {dv.sqref}")
    # 前 12 行内容
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12)):
        vals = [str(c.value)[:22] if c.value is not None else "" for c in row]
        if any(vals):
            print(f"  r{row[0].row}: {vals}")
    if ws.max_row > 12:
        print(f"  ... (共 {ws.max_row} 行)")
