# -*- coding: utf-8 -*-
"""
v3: 修复单元格类型——经纬度/颜色按官方模板写成【文本字符串】(data_type='s')。
末尾附一行国内对照组（若海外被拒而对照组成功→确认范围限制；若全成功→删掉测试行即可）。
"""
import json
import shutil
import warnings
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment

warnings.filterwarnings("ignore")

import _common
BASE = _common.base_from_argv()  # --dir <collection目录> 可指定，缺省取最新
TPL = BASE / "导入模板.xlsx"
OUT = BASE / "amap-import-v3.xlsx"
GEO = BASE / "osm_geocoded_final.json"

COLOR_MAP = {"景点": 1, "美食": 6, "按摩": 8, "住宿": 4, "购物": 5, "交通": 3}


def clip(s, n):
    s = str(s or "")
    return s if len(s) <= n else s[: n - 1] + "…"


def main():
    data = [d for d in json.loads(GEO.read_text(encoding="utf-8")) if d.get("status") == "ok"]

    shutil.copy(TPL, OUT)
    wb = load_workbook(OUT)
    ws = wb["标记位置"]
    ws.delete_rows(2, ws.max_row - 1)

    city_order = {"胡志明市": 0, "芽庄": 1, "岘港": 2}
    cat_order = {"景点": 0, "美食": 1, "按摩": 2, "住宿": 3, "购物": 4, "交通": 5}
    data.sort(key=lambda d: (city_order.get(d["city"], 9), cat_order.get(d["category"], 9)))

    def write_row(vals):
        ws.append(vals)
        r = ws.max_row
        for c in ws[r]:
            c.alignment = Alignment(horizontal="left")

    n = 0
    for d in data:
        addr = " · ".join(["越南", d["city"]] + ([d["address"]] if d.get("address") else []))
        desc = " | ".join(f"《{x}》" for x in d.get("notes", []))
        if d.get("approx"):
            desc += f"｜⚠位置约为{d.get('osm_name', '')}附近，建议手动微调"
        if d.get("address"):
            desc += f"｜{d['address']}"
        write_row([
            clip(d["display"], 40),          # 名称 str
            f"{float(d['lon']):.6f}",        # *经度 -> 字符串!
            f"{float(d['lat']):.6f}",        # *纬度 -> 字符串!
            clip(addr, 120),                 # *地址 str
            str(COLOR_MAP.get(d.get("category"), 7)),  # 颜色 -> 字符串!
            None, None,                      # 图标留空
            clip(desc, 600),                 # 描述 str
            d["city"],                       # 文件夹 str
        ])
        n += 1

    # 对照组: 国内坐标（字符串），用于区分"格式问题"还是"海外范围限制"
    write_row(["测试行-天安门(可删除)", "116.397428", "39.909230", "北京市东城区天安门广场", "1",
               None, None, "对照组：若仅此行成功而其他失败=海外范围限制", "测试"])

    wb.save(OUT)
    print(f"生成 {OUT.name}: {n} 数据行 + 1 对照行")

    # 验证类型
    wb2 = load_workbook(OUT)
    ws2 = wb2["标记位置"]
    for coord in ("B2", "C2", "E2"):
        c = ws2[coord]
        print(f"  {coord} = {c.value!r} type={c.data_type!r}")
    b_last = ws2[f"B{ws2.max_row}"]
    print(f"  对照行 B{ws2.max_row} = {b_last.value!r} type={b_last.data_type!r}")
    print(f"  总行数: {ws2.max_row}")


if __name__ == "__main__":
    main()
